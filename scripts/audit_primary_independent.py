"""Read-only independent checks; never imports the production builder."""
import csv
import json
import math
import argparse
from collections import defaultdict
from pathlib import Path

import openpyxl

BASE = Path(__file__).parent
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--sources', type=Path, default=BASE / 'sources')
parser.add_argument('--data', type=Path, default=BASE.parent / 'data')
parser.add_argument('--report', type=Path)
args = parser.parse_args()
books = {
    p.name: openpyxl.load_workbook(p, read_only=True, data_only=False)
    for p in args.sources.glob('*.xlsx')
}
data = json.loads((args.data / 'primary-data.json').read_text())
records = []
for name in ('primary-observations.csv', 'primary-live-fixed-observations.csv'):
    with (args.data / name).open() as f:
        records.extend(csv.DictReader(f))

groups = defaultdict(list)
seen = set()
for r in records:
    ws = books[r['source_file']][r['source_sheet']]
    identity = (r['source_file'], r['source_sheet'], r['fret_cell'])
    assert identity not in seen, identity
    seen.add(identity)
    value = ws[r['fret_cell']].value
    assert type(value) in (float, int), identity
    assert value == float(r['fret_percent']), identity
    assert ws[r['length_cell']].value == int(r['linker_aa']), identity
    length_col = ''.join(c for c in r['length_cell'] if c.isalpha())
    fret_col = ''.join(c for c in r['fret_cell'] if c.isalpha())
    assert ws[f'{length_col}3'].value == 'Linker Length (AAs)'
    assert ws[f'{fret_col}3'].value == 'FRET Efficiency (%)'
    if r['dataset'] == 'environment':
        assert (length_col, fret_col) in [('C', 'D'), ('I', 'J')]
        assert ws[f'{length_col}2'].value == 'measurements'
        assert r['condition'] == {'Figure 1G': 'vitro', 'Figure 1H': 'cell'}[ws.title]
    else:
        assert (length_col, fret_col) in [('C', 'D'), ('F', 'G')]
        assert ws[f'{length_col}2'].value == {'fixed_23C': 'Fixed (23C)', 'live_37C': 'Live (37C)'}[r['condition']]
    groups[(r['dataset'], r['family'], int(r['linker_aa']), r['condition'])].append(value)

# Independently count all numeric source observations in the measured blocks.
expected_cells = set()
main = books['elife-33927-fig1-data1-v1.xlsx']
supp = books['elife-33927-fig1-figsupp1-data1-v1.xlsx']
for bookname, ws, pairs in [
    ('elife-33927-fig1-data1-v1.xlsx', main['Figure 1G'], [(3, 4), (9, 10)]),
    ('elife-33927-fig1-data1-v1.xlsx', main['Figure 1H'], [(3, 4), (9, 10)]),
    ('elife-33927-fig1-figsupp1-data1-v1.xlsx', supp['Figure 1-Figure supplement 1B'], [(3, 4), (6, 7)]),
]:
    for row in ws.iter_rows(min_row=4):
        for xcol, ycol in pairs:
            x, y = row[xcol - 1], row[ycol - 1]
            if type(x.value) in (float, int) and type(y.value) in (float, int):
                expected_cells.add((bookname, ws.title, y.coordinate))
assert expected_cells == seen

def arithmetic(values):
    n = len(values)
    if not n:
        return 0, None, None
    mean = sum(values) / n
    sd = math.sqrt(sum((x - mean) ** 2 for x in values) / (n - 1)) if n > 1 else None
    return n, mean, sd

def close(a, b):
    assert (a is None and b is None) or math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12), (a, b)

matched = {}
for c in data['constructs']:
    means = {}
    for context in ('vitro', 'cell'):
        raw = groups[('environment', c['family'], c['linker_aa'], context)]
        assert raw == c[context]
        n, mean, sd = arithmetic(raw)
        published = c[f'{context}_summary']
        assert published['n'] == n
        close(mean, published['mean'])
        close(sd, published['sd'])
        means[context] = mean
    if all(x is not None for x in means.values()):
        delta = means['cell'] - means['vitro']
        close(delta, c['shift_pp'])
        matched[c['id']] = delta

controls = {}
for c in data['live_fixed_controls']:
    means = {}
    for label, condition in [('fixed', 'fixed_23C'), ('live', 'live_37C')]:
        raw = groups[('live_fixed', c['family'], c['linker_aa'], condition)]
        assert raw == c[label]
        n, mean, sd = arithmetic(raw)
        assert n == c[f'{label}_summary']['n']
        close(mean, c[f'{label}_summary']['mean'])
        close(sd, c[f'{label}_summary']['sd'])
        means[label] = mean
    controls[c['id']] = means['live'] - means['fixed']
    close(controls[c['id']], c['shift_pp'])

counts = {label: sum(len(v) for k, v in groups.items() if k[3] == label)
          for label in ('vitro', 'cell', 'live_37C', 'fixed_23C')}
assert counts['vitro'] == 111 and counts['cell'] == 739
assert counts['live_37C'] + counts['fixed_23C'] == 425
assert len(data['constructs']) == 16 and len(matched) == 8
outside = {x: sum(abs(v) > x for v in matched.values()) for x in (3, 7, 8)}
assert outside == {3: 7, 7: 4, 8: 0}
for row in data['sensitivity']:
    t = row['tolerance_pp']
    assert set(row['outside']) == {k for k, v in matched.items() if abs(v) > t}
    assert set(row['within']) == {k for k, v in matched.items() if abs(v) <= t}
report = {'status': 'passed', 'source_observations_checked': len(seen),
                  'counts': counts, 'matched_shifts_pp': matched,
                  'outside_tolerance': outside, 'live_fixed_shifts_pp': controls}
if args.report:
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2) + '\n')
print(json.dumps(report, indent=2))
