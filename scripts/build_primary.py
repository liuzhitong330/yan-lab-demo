#!/usr/bin/env python3
"""Rebuild descriptive probe-validation data from two pinned public source workbooks.

Python >=3.10, openpyxl >=3.1. No fitted/model values or inferred force values.
Run: python build_primary.py [--output package] [--cache sources]
"""
from pathlib import Path
import argparse
import csv
import hashlib
import io
import json
import math
import statistics
import urllib.request
import zipfile
from collections import defaultdict
import openpyxl

BASE = "https://cdn.elifesciences.org/articles/33927/"
FILES = {
    "fig1": ("elife-33927-fig1-data1-v1.xlsx", "3c79aa488650280e75071aa177ba9aa73704c011c7014d6c514c0932ea5245a4"),
    "control": ("elife-33927-fig1-figsupp1-data1-v1.xlsx", "4e4cf81e90edb4536a0cbebafd077d635aa48c5b50e51d6531b19685fa670f7e"),
}


def load_source(key, cache):
    filename, expected = FILES[key]
    path = cache / filename
    raw = path.read_bytes() if path.exists() else urllib.request.urlopen(BASE + filename, timeout=60).read()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != expected or not zipfile.is_zipfile(io.BytesIO(raw)):
        raise ValueError(f"Source changed or is not a valid workbook: {filename}, sha256={digest}")
    if not path.exists():
        path.write_bytes(raw)
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=False)


def extract(sheet, xcol, ycol, family, condition, source_key, dataset):
    """Preserve each numeric measured cell and its exact source location."""
    assert sheet.cell(3, xcol).value == "Linker Length (AAs)"
    assert sheet.cell(3, ycol).value == "FRET Efficiency (%)"
    rows = []
    for rownum in range(4, sheet.max_row + 1):
        x = sheet.cell(rownum, xcol).value
        y = sheet.cell(rownum, ycol).value
        if x is None and y is None:
            continue
        if not isinstance(x, (float, int)) or not isinstance(y, (float, int)):
            raise ValueError(f"Non-numeric measured value at {sheet.title}:{rownum}")
        if not math.isfinite(x) or not math.isfinite(y) or int(x) != x or not 0 <= y <= 100:
            raise ValueError(f"Out-of-range source measurement at {sheet.title}:{rownum}")
        if int(x) % len(family):
            raise ValueError("This measured length is not a complete repeat motif")
        rows.append({
            "dataset": dataset, "family": family, "linker_aa": int(x),
            "repeat_count": int(x) // len(family), "condition": condition,
            "fret_percent": y, "source_file": FILES[source_key][0],
            "source_sheet": sheet.title, "source_row": rownum,
            "length_cell": sheet.cell(rownum, xcol).coordinate,
            "fret_cell": sheet.cell(rownum, ycol).coordinate,
        })
    return rows


def stats(values):
    if not values:
        return {"n": 0, "mean": None, "sd": None, "median": None, "min": None, "max": None}
    return {"n": len(values), "mean": statistics.mean(values),
            "sd": statistics.stdev(values) if len(values) > 1 else None,
            "median": statistics.median(values), "min": min(values), "max": max(values)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent / "package")
    parser.add_argument("--cache", type=Path, default=Path(__file__).resolve().parent / "sources")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    args.cache.mkdir(parents=True, exist_ok=True)
    mainbook = load_source("fig1", args.cache)
    controlbook = load_source("control", args.cache)
    raw = []
    for sheetname, condition in [("Figure 1G", "vitro"), ("Figure 1H", "cell")]:
        sheet = mainbook[sheetname]
        for family, xcol, ycol in [("GPGGA", 3, 4), ("GGSGGS", 9, 10)]:
            assert sheet.cell(2, xcol).value == "measurements"
            raw.extend(extract(sheet, xcol, ycol, family, condition, "fig1", "environment"))
    supplemental = controlbook["Figure 1-Figure supplement 1B"]
    assert supplemental.cell(2, 3).value == "Fixed (23C)"
    assert supplemental.cell(2, 6).value == "Live (37C)"
    controls = extract(supplemental, 3, 4, "GGSGGS", "fixed_23C", "control", "live_fixed")
    controls += extract(supplemental, 6, 7, "GGSGGS", "live_37C", "control", "live_fixed")
    grouped = defaultdict(lambda: defaultdict(list))
    for r in raw:
        grouped[(r["family"], r["linker_aa"])][r["condition"]].append(r["fret_percent"])
    constructs = []
    for (family, length), conditions in sorted(grouped.items()):
        vitro, cell = conditions.get("vitro", []), conditions.get("cell", [])
        vstats, cstats = stats(vitro), stats(cell)
        shift = cstats["mean"] - vstats["mean"] if vitro and cell else None
        constructs.append({
            "id": f"{family}-{length}", "family": family, "linker_aa": length,
            "repeat_count": length // len(family), "repeat_segment": family * (length // len(family)),
            "protein_pair": "minimal Clover–mRuby2", "vitro": vitro, "cell": cell,
            "vitro_summary": vstats, "cell_summary": cstats,
            "shift_pp": shift, "abs_shift_pp": abs(shift) if shift is not None else None,
            "notes": "Canonical repeated amino-acid segment, not a complete verified plasmid/insert sequence. Missing environment values are not imputed.",
        })
    controlgroups = defaultdict(lambda: defaultdict(list))
    for r in controls:
        controlgroups[r["linker_aa"]][r["condition"]].append(r["fret_percent"])
    live_fixed = []
    for length, conditions in sorted(controlgroups.items()):
        fixed, live = conditions["fixed_23C"], conditions["live_37C"]
        live_fixed.append({"id": f"GGSGGS-{length}", "family": "GGSGGS", "linker_aa": length,
                           "fixed": fixed, "live": live,
                           "fixed_summary": stats(fixed), "live_summary": stats(live),
                           "shift_pp": statistics.mean(live) - statistics.mean(fixed)})
    matched = [c for c in constructs if c["shift_pp"] is not None]
    sensitivity = [{"tolerance_pp": t, "within": [c["id"] for c in matched if abs(c["shift_pp"]) <= t],
                    "outside": [c["id"] for c in matched if abs(c["shift_pp"]) > t]}
                   for t in [1, 2, 3, 5, 7, 8, 10]]
    # Explicit source-coverage guards prevent accidental model-column ingestion.
    assert len(raw) == 850 and len(controls) == 425
    assert len(constructs) == 16 and len(matched) == 8 and len(live_fixed) == 3
    assert sum(len(c["vitro"]) for c in constructs) == 111
    assert sum(len(c["cell"]) for c in constructs) == 739
    result = {
        "schema_version": "1.0", "title": "Unloaded-FRET transfer check",
        "source": {
            "citation": "LaCroix AS, Lynch AD, Berginski ME, Hoffman BD. eLife 2018;7:e33927.",
            "doi": "https://doi.org/10.7554/eLife.33927", "article": "https://pmc.ncbi.nlm.nih.gov/articles/PMC6053308/",
            "license": "CC BY 4.0", "license_url": "https://creativecommons.org/licenses/by/4.0/",
            "ownership": "External Hoffman Lab benchmark; not Yan Lab probe data.",
            "files": [{"key": k, "url": BASE + v[0], "filename": v[0], "sha256": v[1]} for k, v in FILES.items()],
        },
        "units": {"fret": "%", "shift": "percentage points", "length": "amino acids"},
        "conditions": {
            "vitro": {"label": "In vitro (lysate)", "unit": "experiment measurement", "context": "HEK293-cell hypotonic lysate; spectrofluorometric FRET"},
            "cell": {"label": "In cellulo", "unit": "cell", "context": "Vinculin-null mouse embryonic fibroblasts; fluorescence imaging; pooled from three experiments"},
        },
        "constructs": constructs, "live_fixed_controls": live_fixed, "sensitivity": sensitivity,
        "defaults": {"construct_id": "GPGGA-30", "contrast_id": "GPGGA-10", "tolerance_pp": 3},
        "coverage": {"primary_rows": len(raw), "vitro_rows": 111, "cell_rows": 739, "constructs": 16,
                     "matched_constructs": len(matched), "control_rows": len(controls), "control_comparisons": 3},
        "methods": {
            "shift": "Arithmetic mean in-cell FRET (%) minus arithmetic mean in-vitro FRET (%). Groups are unpaired.",
            "dispersion": "Sample SD uses n−1. SD describes observed dispersion, not standard error, instrument noise, or a confidence interval.",
            "tolerance": "User-chosen absolute mean-shift tolerance. Outside means prioritize a matched-context baseline check; inside does not establish equivalence or calibration validity.",
            "replication": "In-cell data are pooled cells from three experiments; batch identities are not deposited here. No replicate-level confidence intervals, p-values, or independent-cell significance are computed.",
            "comparison": "Environment and assay format/cell context change together. This is a transfer-risk screen, not a causal environment-only estimate.",
            "control": "Supplement1B compares fixed23°C versus live37°C. Temperature and fixation both differ; no pairing or isolated fixation effect is assumed.",
        },
        "limits": [
            "All primary measurements are unloaded FRET. They do not define pN sensitivity, a force-response curve, useful force range, or the best force sensor.",
            "Author-fit columns, model-parameter sheet, and digitized/simulated force datasets are excluded.",
            "Repeat sequences describe the nominal repeated segment only; fluorophore boundaries, junctions, complete insert sequence, and plasmid map require separate sequence verification.",
            "These soluble TSMods and cell contexts do not validate a new targeted Yan Lab sensor. Local force calibration and functional validation remain necessary.",
            "Raw brightness/intensity data for Figure1 supplement1C are not present in the downloaded workbook; no expression-dependence calculation is presented.",
        ],
    }
    (args.output / "primary-data.json").write_text(json.dumps(result, indent=2) + "\n")
    for filename, rows in [("primary-observations.csv", raw), ("primary-live-fixed-observations.csv", controls)]:
        with (args.output / filename).open("w", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    print(json.dumps({"coverage": result["coverage"], "comparisons": [{"id": c["id"], "shift_pp": c["shift_pp"]} for c in matched], "sensitivity": sensitivity}, indent=2))


if __name__ == "__main__":
    main()
